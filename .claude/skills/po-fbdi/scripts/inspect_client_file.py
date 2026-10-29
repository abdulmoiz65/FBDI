# -*- coding: utf-8 -*-
"""
Profiles ANY client-supplied purchase order file and proposes a column mapping.

This is step 1 of the skill. It never guesses silently: every column gets a
confidence and the reason it matched, and everything it could not resolve is
listed separately for the model to decide.

  python inspect_client_file.py "inbox/whatever the client sent.xlsx"
  python inspect_client_file.py file.xlsx --sheet "PO Data" --json profile.json

Output: a human-readable profile on stdout and, with --json, the same data as
JSON. The JSON is what you edit into a mapping file for convert_to_fbdi.py.
"""
import argparse
import csv
import datetime
import difflib
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import po_columns as P  # noqa: E402

# captions that mean "this column is not data"
JUNK = re.compile(r"^(unnamed|column\d*|field\d*|col\d*|#|s\.?no\.?|sr\.?no\.?|serial)", re.I)


def slug(s):
    """Lower-case, strip punctuation and collapse spaces, for header matching."""
    if s is None:
        return ""
    s = str(s).replace("*", " ")
    s = re.sub(r"[‐-―]", "-", s)          # unicode dashes
    s = re.sub(r"[^0-9a-zA-Z]+", " ", s)
    return re.sub(r"\s+", " ", s).strip().lower()


# canonical name and synonyms, pre-slugged, in one lookup
LOOKUP = {}
for _c in P.COLUMNS:
    LOOKUP.setdefault(slug(_c["name"]), _c["key"])
for _k, _alts in P.SYNONYMS.items():
    for _a in _alts:
        LOOKUP.setdefault(slug(_a), _k)


def looks_like_header(values):
    """Score a row on how much it looks like a header row."""
    vals = [v for v in values if v is not None and str(v).strip() != ""]
    if len(vals) < 3:
        return 0.0
    strings = [v for v in vals if isinstance(v, str)]
    if len(strings) < len(vals) * 0.7:
        return 0.0
    short = [s for s in strings if len(s) <= 60 and "\n" not in s.strip()]
    known = sum(1 for s in strings if slug(s) in LOOKUP)
    uniq = len({slug(s) for s in strings})
    return (len(short) / len(vals)) * 0.4 + (uniq / len(vals)) * 0.2 + \
           min(known / max(len(vals), 1), 1.0) * 1.4


def read_grid(path, sheet=None, limit=400):
    """Return (sheet_name, all sheet names, grid of rows) for xlsx/xlsm/csv."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            rows = [r for _i, r in zip(range(limit), csv.reader(fh))]
        return os.path.basename(path), [os.path.basename(path)], rows

    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    if sheet and sheet in names:
        ws = wb[sheet]
    else:
        # the sheet with the most data wins, ignoring obvious instruction sheets
        best, ws = -1, wb.worksheets[0]
        for cand in wb.worksheets:
            if cand.sheet_state != "visible":
                continue
            score = (cand.max_row or 0) * min(cand.max_column or 0, 80)
            if re.search(r"overview|instruction|revision|history|read ?me|lists?",
                         cand.title, re.I):
                score *= 0.05
            if score > best:
                best, ws = score, cand
    grid = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, limit)):
        grid.append([c.value for c in row])
    return ws.title, names, grid


def detect_header(grid):
    best_row, best_score = 0, -1.0
    for i, row in enumerate(grid[:20]):
        s = looks_like_header(row)
        if s > best_score:
            best_row, best_score = i, s
    return best_row, round(best_score, 3)


def match_column(caption):
    """Best canonical key for one client caption. Returns (key, confidence, how)."""
    s = slug(caption)
    if not s or JUNK.match(s):
        return None, 0.0, "ignored"
    if s in LOOKUP:
        return LOOKUP[s], 1.0, "exact"
    # containment: "po line number (legacy)" contains "po line number"
    for k in sorted((k for k in LOOKUP if len(k) > 3), key=len, reverse=True):
        if (" " + k + " ") in (" " + s + " "):
            return LOOKUP[k], 0.92, "contains '%s'" % k
    hit = difflib.get_close_matches(s, list(LOOKUP), n=1, cutoff=0.86)
    if hit:
        return LOOKUP[hit[0]], round(difflib.SequenceMatcher(None, s, hit[0]).ratio(), 2), \
            "fuzzy '%s'" % hit[0]
    return None, 0.0, "no match"


def resolve_conflicts(matches):
    """Two client columns can both look like one field. Never pick silently.

    matches is [(index, key, conf, how)]. The strongest candidate keeps the
    field; every rival is unmapped and told which field it competed for, so the
    caller has to make the call. Classic case: 'Qty Ordered' vs 'Balance Qty'.
    """
    by_key, conflicts = {}, {}
    for i, key, conf, _how in matches:
        if key:
            by_key.setdefault(key, []).append((conf, -i, i))
    for key, cands in by_key.items():
        if len(cands) < 2:
            continue
        cands.sort(reverse=True)
        winner = cands[0][2]
        for _c, _n, i in cands[1:]:
            conflicts[i] = key
        conflicts.setdefault(winner, None)
    return conflicts


def profile_column(rows, idx):
    vals = [r[idx] if idx < len(r) else None for r in rows]
    filled = [v for v in vals if v is not None and str(v).strip() != ""]
    samples, seen = [], set()
    for v in filled:
        t = str(v).strip()
        if isinstance(v, (datetime.datetime, datetime.date)):
            t = v.strftime("%Y-%m-%d")
        if t not in seen:
            seen.add(t)
            samples.append(t[:40])
        if len(samples) >= 4:
            break
    kinds = set()
    for v in filled[:60]:
        if isinstance(v, (datetime.datetime, datetime.date)):
            kinds.add("date")
        elif isinstance(v, (int, float)):
            kinds.add("number")
        else:
            kinds.add("text")
    return {
        "fill": round(len(filled) / max(len(vals), 1), 2),
        "distinct": len({str(v).strip() for v in filled}),
        "kind": "/".join(sorted(kinds)) or "empty",
        "samples": samples,
    }


def main():
    ap = argparse.ArgumentParser(description="Profile a client PO file and propose a mapping.")
    ap.add_argument("infile")
    ap.add_argument("--sheet", default=None, help="force a sheet name")
    ap.add_argument("--header-row", type=int, default=None, help="1-based header row")
    ap.add_argument("--json", default=None, help="write the profile to this JSON file")
    a = ap.parse_args()

    sheet, all_sheets, grid = read_grid(a.infile, a.sheet)
    if not grid:
        raise SystemExit("The file appears to be empty.")

    if a.header_row:
        hrow, hscore = a.header_row - 1, 1.0
    else:
        hrow, hscore = detect_header(grid)

    captions = grid[hrow]
    body = [r for r in grid[hrow + 1:]
            if any(v is not None and str(v).strip() != "" for v in r)]

    # our own template is identifiable by its Row Check column, not by a
    # caption count - a well-named client export also scores high on that
    ours = any(slug(c) == "row check" for c in captions)
    skipped = 0
    if ours and len(body) >= 3:
        # guidance row, data-type row, then the yellow example row
        marker = " ".join(str(v) for v in body[2][:2] if v)
        skipped = 3 if "EXAMPLE" in marker.upper() else 2
        body = body[skipped:]

    raw = []
    for i in range(len(captions)):
        key, conf, how = match_column(captions[i])
        raw.append((i, key, conf, how))
    conflicts = resolve_conflicts(raw)

    used, columns = set(), []
    for i, key, conf, how in raw:
        prof = profile_column(body, i)
        if i in conflicts and conflicts[i] is not None:
            rival = P.by_key(conflicts[i])["name"]
            key, conf, how = None, 0.0, "also looks like '%s' - you decide" % rival
        if key and prof["fill"] == 0.0 and conf < 1.0:
            key, conf, how = None, 0.0, "empty column, match not trusted"
        if key:
            used.add(key)
        col = P.by_key(key) if key else None
        columns.append({
            "index": i + 1,
            "letter": get_column_letter(i + 1),
            "caption": None if captions[i] is None else str(captions[i]).strip(),
            "key": key,
            "maps_to": col["name"] if col else None,
            "confidence": conf,
            "how": how,
            **prof,
        })

    required = [k for keys in P.MANDATORY_GROUPS.values() for k in keys]
    missing = [P.by_key(k)["name"] for k in required if k not in used]
    review = [c for c in columns if c["key"] and c["confidence"] < 1.0]
    unmapped = [c for c in columns if not c["key"] and c["fill"] > 0]

    out = {
        "file": os.path.abspath(a.infile),
        "sheet": sheet,
        "all_sheets": all_sheets,
        "header_row": hrow + 1,
        "header_confidence": hscore,
        "data_starts_row": hrow + 2 + skipped,
        "data_rows": len(body),
        "is_standard_template": bool(ours),
        "columns": columns,
        "mapped": sorted(used),
        "missing_mandatory": missing,
        "needs_review": [c["caption"] for c in review],
        "unmapped_with_data": [c["caption"] for c in unmapped],
    }

    print("file            : %s" % out["file"])
    print("sheet           : %s   (sheets: %s)" % (sheet, ", ".join(all_sheets)))
    print("header row      : %d   (confidence %.2f)" % (out["header_row"], hscore))
    print("data starts row : %d   (%d data rows)" % (out["data_starts_row"], len(body)))
    print("standard template: %s" % ("yes" if ours else "no - client's own layout"))
    print()
    print("%-4s %-34s %-28s %-6s %-5s %s" %
          ("COL", "CLIENT CAPTION", "MAPS TO", "CONF", "FILL", "SAMPLES"))
    print("-" * 118)
    for c in columns:
        if c["caption"] is None and c["fill"] == 0:
            continue
        print("%-4s %-34s %-28s %-6s %-5s %s" % (
            c["letter"], (c["caption"] or "")[:34], (c["maps_to"] or "-")[:28],
            c["confidence"] or "", c["fill"], " | ".join(c["samples"])[:46]))

    if review:
        print("\nCHECK THESE MATCHES (confidence below 1.0):")
        for c in review:
            print("  %s  '%s' -> %s   [%s]" % (c["letter"], c["caption"], c["maps_to"], c["how"]))
    if unmapped:
        print("\nNOT MAPPED BUT HOLDS DATA:")
        for c in unmapped:
            print("  %s  '%s'  [%s]  samples: %s" % (
                c["letter"], c["caption"], c["how"], " | ".join(c["samples"])[:52]))
    if missing:
        print("\nMANDATORY FIELDS WITH NO SOURCE COLUMN:")
        for m in missing:
            print("  - %s" % m)
        print("  These must come from a constant (same for every PO) or be asked for.")

    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump(out, fh, indent=1, ensure_ascii=False)
        print("\nprofile written to %s" % a.json)


if __name__ == "__main__":
    main()
