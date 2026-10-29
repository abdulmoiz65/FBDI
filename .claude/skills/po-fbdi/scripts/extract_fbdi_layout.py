# -*- coding: utf-8 -*-
"""
Regenerates assets/fbdi_layout.json from Oracle's own FBDI template.

The layout asset is the column order of each interface, and the converter writes
CSV fields in exactly that order. So it has to be the REAL column list.

Why this script exists: the asset was first produced ad hoc from each sheet's
`max_column`, which on PO_LINE_LOCATIONS_INTERFACE is 125 even though Oracle
only captions 97 columns - the remaining 28 carry stray formatting and no
header. Every generated row was padded with 28 phantom trailing fields. Reading
to the last CAPTIONED cell instead is the fix, and keeping it in a script means
the next Oracle template drop can be re-extracted the same way.

Run:  python .claude/skills/po-fbdi/scripts/extract_fbdi_layout.py
      python .claude/skills/po-fbdi/scripts/extract_fbdi_layout.py --check
"""
import argparse
import io
import json
import os
import sys

import openpyxl

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(SCRIPTS)
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(SKILL)))
TEMPLATE = os.path.join(ROOT, "reference", "POPurchaseOrderImportTemplate.xlsm")
OUT = os.path.join(SKILL, "assets", "fbdi_layout.json")

SHEETS = ["PO_HEADERS_INTERFACE", "PO_LINES_INTERFACE",
          "PO_LINE_LOCATIONS_INTERFACE", "PO_DISTRIBUTIONS_INTERFACE"]
HEADER_ROW = 4
DATA_ROW = 5
ZIP_NAME = "PoImportOrders.zip"


def extract(path=TEMPLATE):
    wb = openpyxl.load_workbook(path, keep_vba=True, read_only=True)
    sheets = {}
    for name in SHEETS:
        ws = wb[name]
        caps = [str(ws.cell(HEADER_ROW, c).value or "").strip()
                for c in range(1, ws.max_column + 1)]
        # stop at the last cell that actually carries a caption
        last = max((i for i, v in enumerate(caps, start=1) if v), default=0)
        cols = caps[:last]
        blanks = [i for i, v in enumerate(cols, start=1) if not v]
        if blanks:
            raise SystemExit(
                "%s has un-captioned columns before the last caption, at %s. "
                "The column order cannot be trusted - inspect the template."
                % (name, blanks))
        sheets[name] = cols
    return {
        "source": os.path.basename(path),
        "zip_name": ZIP_NAME,
        "header_row": HEADER_ROW,
        "data_row": DATA_ROW,
        "sheets": sheets,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true",
                    help="compare the asset with the template and exit non-zero "
                         "if they differ, instead of rewriting it")
    a = ap.parse_args()

    fresh = extract()
    for name, cols in fresh["sheets"].items():
        print("%-30s %3d columns" % (name, len(cols)))

    if a.check:
        with io.open(OUT, encoding="utf-8") as fh:
            current = json.load(fh)
        drift = [n for n in SHEETS
                 if current["sheets"].get(n) != fresh["sheets"][n]]
        if drift:
            for n in drift:
                print("  DRIFT %s: asset has %d, template has %d"
                      % (n, len(current["sheets"].get(n, [])), len(fresh["sheets"][n])))
            sys.exit(1)
        print("asset matches the template")
        return

    with io.open(OUT, "w", encoding="utf-8") as fh:
        json.dump(fresh, fh, indent=1, ensure_ascii=False)
    print("written: %s" % OUT)


if __name__ == "__main__":
    main()
