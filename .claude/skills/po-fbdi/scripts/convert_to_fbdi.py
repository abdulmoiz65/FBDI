# -*- coding: utf-8 -*-
"""
Converts a filled "Standard - Purchase Order Data Template.xlsx" into the Oracle
Purchase Order Import FBDI deliverables.

Outputs (into --outdir, default: output/<timestamp-free folder name>):
  PO_HEADERS_INTERFACE.csv
  PO_LINES_INTERFACE.csv
  PO_LINE_LOCATIONS_INTERFACE.csv
  PO_DISTRIBUTIONS_INTERFACE.csv
  PoImportOrders.zip                     <- upload this in "Load Interface File for Import"
  POPurchaseOrderImportTemplate_filled.xlsm   (optional, --xlsm) for review in Oracle's own template
  validation_report.txt

Usage:
  python .claude/skills/po-fbdi/scripts/convert_to_fbdi.py "inbox/Client.xlsx"
  python .claude/skills/po-fbdi/scripts/convert_to_fbdi.py "inbox/Client.xlsx" \n         --map "inbox/Client.mapping.json" --batch 9100 --xlsm
"""
import argparse
import csv
import datetime
import json
import os
import re
import shutil
import sys
import zipfile

import openpyxl
from openpyxl.utils import column_index_from_string

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import po_columns as P  # noqa: E402

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(SCRIPTS)                       # .claude/skills/po-fbdi
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(SKILL)))
ORACLE_TEMPLATE = os.path.join(ROOT, "reference", "POPurchaseOrderImportTemplate.xlsm")
LAYOUT_JSON = os.path.join(SKILL, "assets", "fbdi_layout.json")

SHEET_ORDER = ["PO_HEADERS_INTERFACE", "PO_LINES_INTERFACE",
               "PO_LINE_LOCATIONS_INTERFACE", "PO_DISTRIBUTIONS_INTERFACE"]

DATE_OUT = "%Y/%m/%d"

# row numbers inside Oracle's own FBDI template
ORA_HEADER_ROW = 4
ORA_DATA_ROW = 5


# --------------------------------------------------------------- small helpers
def norm(s):
    """Normalise an FBDI column caption so it can be matched reliably."""
    if s is None:
        return ""
    s = str(s).replace("*", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip().lower()


def txt(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime(DATE_OUT)
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def as_date(v):
    """Accept a real date, or a YYYY-MM-DD / DD-MM-YYYY / YYYY/MM/DD string."""
    if v is None or v == "":
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime(DATE_OUT)
    s = str(v).strip().replace(".", "-").replace("/", "-")
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y", "%d-%b-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(s, f).strftime(DATE_OUT)
        except ValueError:
            continue
    return str(v).strip()          # hand it over unchanged and flag it


def as_num(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "").replace(" ", "")
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except ValueError:
        return None                # signals "not a number"


def open_workbook(path):
    """Open .xlsx / .xlsm, or a .csv presented as a one-sheet workbook.

    Clients do send CSV extracts, and inspect_client_file.py already profiles
    them, so refusing them here would be an inconsistency. Every CSV value
    arrives as text; as_date() and as_num() already cope with that.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = os.path.basename(path)
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            for row in csv.reader(fh):
                ws.append(row)
        return wb
    if ext not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        raise SystemExit("Cannot read '%s'. Supported: .xlsx, .xlsm, .csv." % ext)
    return openpyxl.load_workbook(path, data_only=True)


# --------------------------------------------------------------- derivation
# What distinguishes one SCHEDULE of a line from another: when it is wanted and
# where it goes. What distinguishes one DISTRIBUTION of a schedule: which
# account / project / person it is charged to.
SCHEDULE_KEYS = P.SCHEDULE_KEYS
DIST_KEYS = P.DIST_KEYS


def sig(rec, keys):
    return tuple(txt(rec.get(k)) for k in keys)


def derive_structure(rows, present):
    """Fill in the schedule / distribution shape the lean sheet does not ask for.

    The lean client sheet has no Schedule No. or Distribution No. at all - a
    Fusion implementation should not make a client learn Oracle's PO structure.
    Instead the client repeats the row and changes what is genuinely different,
    and the numbering is worked out here:

      * rows of one line that differ in date or ship-to  -> separate SCHEDULES
      * rows of one schedule that differ in account etc. -> separate DISTRIBUTIONS
      * a repeated row that differs in nothing at all    -> still a duplicate error

    Quantities roll UP, because that is the direction the client can actually
    fill in: the quantity on a row is that part's quantity, a schedule is the
    sum of its distributions, and the line is the sum of its schedules.

    Returns a list of notes describing every line that was treated as a split,
    so a split created by accident is visible in the report rather than silent.
    """
    notes = []
    have_sched = "schedule_number" in present
    have_dist = "distribution_number" in present

    groups = {}
    for rec in rows:
        groups.setdefault((txt(rec.get("po_number")), txt(rec.get("line_number"))), []).append(rec)

    for (po, line), grp in groups.items():
        # ---- schedule numbering
        if not have_sched:
            order, seen = [], {}
            for rec in grp:
                k = sig(rec, SCHEDULE_KEYS)
                if k not in seen:
                    seen[k] = len(order) + 1
                    order.append(k)
                rec["schedule_number"] = seen[k]
        # ---- distribution numbering, within each schedule
        if not have_dist:
            per_sched = {}
            for rec in grp:
                sn = txt(rec.get("schedule_number")) or "1"
                seen = per_sched.setdefault(sn, {})
                k = sig(rec, DIST_KEYS)
                if k not in seen:
                    seen[k] = len(seen) + 1
                rec["distribution_number"] = seen[k]

        if len(grp) > 1:
            nsched = len({txt(r.get("schedule_number")) for r in grp})
            ndist = len(grp)
            notes.append("PO %s line %s: %d rows read as %d schedule(s) and %d distribution(s). "
                         "Quantities were added up, not repeated."
                         % (po, line, len(grp), nsched, ndist))

    # ---- quantities and amounts roll up from the rows
    # One client column carries both: a quantity on a Goods line, an amount on a
    # service line. Whichever it is, the row holds THAT PART's figure, a schedule
    # is the sum of its distributions, and the line is the sum of its schedules.
    def measure(rec):
        if txt(rec.get("line_type")) == "Goods":
            return as_num(rec.get("quantity"))
        return as_num(rec.get("line_amount"))

    if "distribution_quantity" not in present:
        for rec in rows:
            rec["distribution_quantity"] = measure(rec)

    for (po, line), grp in groups.items():
        goods = txt(grp[0].get("line_type")) == "Goods"
        by_sched = {}
        for rec in grp:
            by_sched.setdefault(txt(rec.get("schedule_number")), []).append(rec)

        line_total, seen_any = 0, False
        for part in by_sched.values():
            vals = [as_num(r.get("distribution_quantity")) for r in part]
            vals = [v for v in vals if isinstance(v, (int, float))]
            sched_total = round(sum(vals), 6) if vals else ""
            if "schedule_quantity" not in present:
                for r in part:
                    r["schedule_quantity"] = sched_total
            if vals:
                line_total += sum(vals)
                seen_any = True

        total = round(line_total, 6) if seen_any else ""
        for rec in grp:
            rec["_line_total_qty"] = total if goods else ""
            rec["_line_total_amount"] = "" if goods else total

    # ---- deliver-to falls back to the ship-to on the same row
    for rec in rows:
        if txt(rec.get("deliver_to_location")) == "":
            rec["deliver_to_location"] = rec.get("ship_to_location")

    return notes


# --------------------------------------------------------------- read template
def read_with_mapping(path, mapping, auto_number=False):
    """Read a client's own layout using an explicit mapping file.

    mapping = {
      "sheet": "Sheet1", "data_starts_row": 4,
      "columns":   {"po_number": "A", "quantity": 7, ...},   letter or 1-based index
      "constants": {"doc_type": "STANDARD", "procurement_bu": "Example BU"},
      "overrides": {"7": {"line_amount": 180000}}            per source-row patches
    }

    Constants fill a field the client file has no column for. Overrides patch
    individual cells, keyed by the row number in the client file, so the file
    we received stays untouched and every correction stays visible in one place.
    """
    wb = open_workbook(path)
    name = mapping.get("sheet")
    ws = wb[name] if name and name in wb.sheetnames else wb.worksheets[0]
    start = int(mapping.get("data_starts_row") or 2)

    pos = {}
    for key, ref in (mapping.get("columns") or {}).items():
        try:
            P.by_key(key)
        except KeyError:
            raise SystemExit("Mapping refers to unknown field '%s'." % key)
        pos[key] = ref if isinstance(ref, int) else column_index_from_string(str(ref).strip().upper())
    consts = mapping.get("constants") or {}
    for key in consts:
        try:
            P.by_key(key)
        except KeyError:
            raise SystemExit("Mapping constant refers to unknown field '%s'." % key)

    if "po_number" not in pos and "po_number" not in consts and not auto_number:
        raise SystemExit("The mapping has no source for 'po_number'.")

    overrides = {int(k): v for k, v in (mapping.get("overrides") or {}).items()}
    for patch in overrides.values():
        for key in patch:
            try:
                P.by_key(key)
            except KeyError:
                raise SystemExit("Mapping override refers to unknown field '%s'." % key)

    rows, applied = [], []
    for r in range(start, ws.max_row + 1):
        rec = {"_row": r}
        for key, c in pos.items():
            rec[key] = P.normalise_value(key, ws.cell(r, c).value)
        for key, v in consts.items():
            if txt(rec.get(key)) == "":
                rec[key] = v
        for key, v in overrides.get(r, {}).items():
            rec[key] = P.normalise_value(key, v)
            applied.append("row %d: %s = %s" % (r, P.by_key(key)["name"], v))
        if txt(rec.get("po_number")) == "" and not (auto_number and not row_is_empty(rec)):
            continue
        rows.append(rec)

    if auto_number:
        assign_auto_po_numbers(rows)

    errors = []
    if not rows:
        errors.append("No data rows found from row %d of sheet '%s'." % (start, ws.title))
    unused = sorted(set(overrides) - {r["_row"] for r in rows})
    if unused:
        errors.append("Overrides target rows that hold no PO: %s."
                      % ", ".join(str(u) for u in unused))
    return rows, errors, applied, set(pos) | set(consts)


def read_client_sheet(path, auto_number=False):
    """Return (rows, errors) for a file that uses our standard template."""
    wb = open_workbook(path)
    ws = wb[P.DATA_SHEET] if P.DATA_SHEET in wb.sheetnames else wb.worksheets[-1]

    # locate the field-name row (survives inserted rows above)
    want = norm(P.COLUMNS[1]["name"])
    name_row = None
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            if norm(ws.cell(r, c).value).startswith(want):
                name_row = r
                break
        if name_row:
            break
    if name_row is None:
        raise SystemExit("Could not find the header row (a cell containing '%s') in sheet '%s'."
                         % (P.COLUMNS[1]["name"], ws.title))

    # map every template column to its actual position by caption
    caption = {}
    for c in range(1, ws.max_column + 1):
        caption[norm(ws.cell(name_row, c).value)] = c
    # Match against the whole catalogue, not the active profile: the file in
    # front of us may be the wide sheet, or an older template, and any column it
    # does carry should be read rather than silently ignored.
    pos, missing = {}, []
    for col in P.CATALOGUE:
        n = norm(col["name"])
        hit = caption.get(n)
        if hit is None:                                   # tolerate the " *" suffix
            for cap, c in caption.items():
                if cap.startswith(n):
                    hit = c
                    break
        if hit is None:
            missing.append(col["name"])
        else:
            pos[col["key"]] = hit

    errors = []
    # A field Oracle derives is expected to be absent from the lean sheet - the
    # converter fills it in. Only a core field going missing is an error.
    core_missing = [n for n in missing
                    if not any(c["name"] == n and c["adv"] for c in P.CATALOGUE)]
    if core_missing:
        errors.append("Columns not found in the client file (ignored): " + ", ".join(core_missing))

    po_col = pos.get("po_number")
    if not po_col:
        raise SystemExit(
            "No 'PO Number' column found, so this file does not use the standard "
            "template. Run inspect_client_file.py on it, build a mapping file, "
            "then re-run with --map mapping.json.")

    rows = []
    for r in range(name_row + 3, ws.max_row + 1):
        marker = ws.cell(r, 1).value
        if isinstance(marker, str) and "EXAMPLE" in marker.upper():
            continue
        blank_po = txt(ws.cell(r, po_col).value) == ""
        if blank_po and not auto_number:
            continue
        rec = {"_row": r}
        for key, c in pos.items():
            rec[key] = P.normalise_value(key, ws.cell(r, c).value)
        if blank_po and row_is_empty(rec):
            continue
        rows.append(rec)
    if auto_number:
        assign_auto_po_numbers(rows)
    return rows, errors, set(pos)


AUTO_PO_NOTES = []


def row_is_empty(rec):
    """True when the client typed nothing on this row.

    A blank PO Number used to be the end-of-data marker. With --auto-number it
    is the normal case, so emptiness has to be judged on the rest of the row:
    every field the sheet carries is blank apart from Row Check, which holds a
    formula (or a stale result someone typed over) on rows the client never
    filled in.
    """
    return not any(txt(v) for k, v in rec.items()
                   if k not in ("_row", "check") and not k.startswith("_"))


def assign_auto_po_numbers(rows):
    """Give rows with no PO number an identity that never reaches the CSV.

    Oracle assigns the document number itself when the Order column is blank,
    which is what the consultant confirmed for this instance. The converter
    still needs a PO identity to group lines, schedules and distributions, so
    one is derived from the PO-level fields: rows whose header values are all
    identical are one PO, rows that differ in any of them are separate POs.
    The synthetic value is stripped again in build_fbdi.
    """
    hdr_keys = [c["key"] for c in P.CATALOGUE
                if c["tgt"] and c["tgt"][0] == "PO_HEADERS_INTERFACE"
                and c["key"] != "po_number"]
    seen = {}
    for rec in rows:
        if txt(rec.get("po_number")) != "":
            continue
        k = tuple(txt(rec.get(x)) for x in hdr_keys)
        if k not in seen:
            seen[k] = "(auto %d)" % (len(seen) + 1)
        rec["po_number"] = seen[k]
        rec["_auto_po"] = True
    if seen:
        AUTO_PO_NOTES.append("%d row(s) carry no PO number, grouped into %d PO(s) [%s] by their "
                     "PO-level fields. The Order column is written BLANK so Oracle assigns "
                     "the number from the instance's document numbering."
                     % (sum(1 for r in rows if r.get("_auto_po")), len(seen),
                        ", ".join(sorted(seen.values()))))


# --------------------------------------------------------------- validation
def validate(rows, expected_segments=None, present=None, derive_ship_to_org=False):
    """Return (errors, warnings). Mirrors the Row Check formula in the template.

    `present` is the set of field keys the client file actually carried. A rule
    about a field Oracle derives must not fire when that field was never asked
    for - otherwise the lean sheet would fail its own validation.
    """
    if present is None:
        present = {c["key"] for c in P.CATALOGUE}

    # The advice in a duplicate-row error has to match the sheet the client was
    # actually given: the lean sheet has no Schedule No. to increase.
    if "schedule_number" in present or "distribution_number" in present:
        split_hint = "Increase the Schedule No. or Distribution No. to split it."
    else:
        split_hint = ("If it was meant to be a separate delivery or a separate account, change the "
                      "Need-by Date, Ship-to Location or Charge Account on it. If it is a copy made "
                      "by mistake, delete it.")

    def have(*keys):
        return all(k in present for k in keys)
    errors, warnings = [], []
    blank_accounts = []
    hdr_keys = [c["key"] for c in P.CATALOGUE
                if c["tgt"] and c["tgt"][0] == "PO_HEADERS_INTERFACE"]

    seen_header = {}
    seen_dist = {}
    for rec in rows:
        r = rec["_row"]
        po = txt(rec.get("po_number"))

        for keys in P.MANDATORY_GROUPS.values():
            for k in keys:
                if txt(rec.get(k)) == "":
                    errors.append("Row %d (PO %s): '%s' is mandatory but empty."
                                  % (r, po, P.by_key(k)["name"]))

        lt = txt(rec.get("line_type"))
        if lt == "Goods":
            for k in ("quantity", "unit_price", "uom"):
                if txt(rec.get(k)) == "":
                    errors.append("Row %d (PO %s): '%s' is mandatory for a Goods line."
                                  % (r, po, P.by_key(k)["name"]))
        elif lt and txt(rec.get("line_amount")) == "":
            errors.append("Row %d (PO %s): 'Line Amount' is mandatory for a service line." % (r, po))
        elif lt and lt not in P.LISTS["LineType"]:
            errors.append("Row %d (PO %s): Line Type '%s' is not one of %s."
                          % (r, po, lt, ", ".join(P.LISTS["LineType"])))

        if txt(rec.get("item")) == "" and txt(rec.get("category")) == "":
            errors.append("Row %d (PO %s): 'Category Name' is mandatory when 'Item Number' is empty." % (r, po))

        dest = txt(rec.get("destination_type")).upper()
        if dest and dest not in P.LISTS["DestType"]:
            errors.append("Row %d (PO %s): Destination Type '%s' is not one of %s."
                          % (r, po, dest, ", ".join(P.LISTS["DestType"])))
        if dest == "INVENTORY" and txt(rec.get("ship_to_org")) == "":
            # Fusion can resolve the inventory organization from the ship-to
            # location, which is assigned to one in Manage Inventory
            # Organizations. That only holds where the setup actually carries
            # the assignment, so it is opt-in per load and reported, never
            # assumed: without the flag a blank org is still an error.
            if derive_ship_to_org:
                warnings.append("Row %d (PO %s): 'Ship-to Organization' left blank for INVENTORY - "
                                "Oracle must derive it from Ship-to Location '%s'. Confirm that "
                                "location is assigned to an inventory organization in the instance, "
                                "or the row rejects."
                                % (r, po, txt(rec.get("ship_to_location"))))
            else:
                errors.append("Row %d (PO %s): 'Ship-to Organization' is mandatory for INVENTORY." % (r, po))
        # Not an error any more: on a Fusion implementation the account
        # generator builds the charge account from item, category, destination
        # and requester. Blank means "let Oracle derive it", which is the normal
        # choice. It is reported once, not per row, so the consultant confirms
        # account generation is configured before the load.
        if dest == "EXPENSE" and txt(rec.get("charge_account")) == "":
            blank_accounts.append(r)

        if have("comm_method", "supplier_email") and                 txt(rec.get("comm_method")).upper() == "EMAIL" and txt(rec.get("supplier_email")) == "":
            errors.append("Row %d (PO %s): 'Supplier Email' is mandatory when Communication Method = EMAIL." % (r, po))

        if have("conversion_rate", "rate_type") and                 txt(rec.get("conversion_rate")) != "" and txt(rec.get("rate_type")) == "":
            errors.append("Row %d (PO %s): 'Rate Type' is mandatory when a Conversion Rate is given "
                          "(use User for a rate you supply yourself)." % (r, po))

        for k in ("quantity", "unit_price", "line_amount", "schedule_quantity",
                  "distribution_quantity", "conversion_rate", "line_number"):
            if txt(rec.get(k)) != "" and as_num(rec.get(k)) is None:
                errors.append("Row %d (PO %s): '%s' = '%s' is not a valid number."
                              % (r, po, P.by_key(k)["name"], txt(rec.get(k))))
        for k in ("need_by_date", "promised_date", "expenditure_date"):
            v = rec.get(k)
            if txt(v) != "" and not isinstance(v, (datetime.date, datetime.datetime)) \
                    and not re.match(r"^\d{4}/\d{2}/\d{2}$", as_date(v)):
                warnings.append("Row %d (PO %s): '%s' = '%s' could not be read as a date - please confirm."
                                % (r, po, P.by_key(k)["name"], txt(v)))

        # header values must be identical on every row of the same PO
        sig = {k: txt(rec.get(k)) for k in hdr_keys if k != "po_number"}
        if po in seen_header:
            first_row, first_sig = seen_header[po]
            for k, v in sig.items():
                if v != first_sig.get(k, ""):
                    errors.append("Row %d: PO %s has '%s' = '%s' but row %d has '%s'. "
                                  "PO level fields must be identical on every line of the same PO."
                                  % (r, po, P.by_key(k)["name"], v, first_row, first_sig.get(k, "")))
        else:
            seen_header[po] = (r, sig)

        # duplicate line / schedule / distribution
        dkey = (po, txt(rec.get("line_number")),
                txt(as_num(rec.get("schedule_number")) or 1),
                txt(as_num(rec.get("distribution_number")) or 1))
        if dkey in seen_dist:
            errors.append("Row %d: PO %s line %s schedule %s distribution %s is already on row %d. %s"
                          % ((r,) + dkey + (seen_dist[dkey], split_hint)))
        else:
            seen_dist[dkey] = r

    if blank_accounts:
        warnings.append(
            "%d EXPENSE row(s) have no Charge Account, so Oracle's account generator will derive it "
            "(rows %s%s). Confirm account generation is configured for this BU, otherwise the import "
            "will reject them."
            % (len(blank_accounts), ", ".join(str(r) for r in blank_accounts[:8]),
               " ..." if len(blank_accounts) > 8 else ""))

    if have("schedule_quantity") or have("distribution_quantity"):
        errors += check_splits(rows)
    errors += check_charge_accounts(rows, expected_segments)
    return errors, warnings


SEGMENT_SPLIT = re.compile(r"[-.|]")


def check_charge_accounts(rows, expected=None):
    """Every charge account must split into the same number of segments.

    The client types one hyphen-delimited string and we fan it out into
    CHARGE_ACCOUNT_SEGMENT1..30. If one row has five segments and another has
    six, the sixth value lands in a different segment for different rows and the
    load posts to the wrong accounts without erroring. Oracle cannot catch that
    for us, so it has to be caught here.

    `expected` comes from "charge_account_segments" in the mapping file. When it
    is set the count is enforced; when it is not, the counts merely have to
    agree with each other, and the caller is told what was inferred.
    """
    errors = []
    seen = {}
    for rec in rows:
        acct = txt(rec.get("charge_account"))
        if not acct:
            continue
        n = len([s for s in SEGMENT_SPLIT.split(acct)])
        seen.setdefault(n, []).append((rec["_row"], acct))

    if not seen:
        return errors

    if expected is not None:
        for n, hits in sorted(seen.items()):
            if n != expected:
                row, acct = hits[0]
                errors.append(
                    "Row %d: charge account '%s' has %d segment(s) but the mapping says the chart of "
                    "accounts has %d. %d row(s) affected. Fix the data or correct "
                    "'charge_account_segments'." % (row, acct, n, expected, len(hits)))
        return errors

    if len(seen) > 1:
        detail = "; ".join(
            "%d segment(s) on %d row(s) e.g. row %d '%s'" % (n, len(h), h[0][0], h[0][1])
            for n, h in sorted(seen.items()))
        errors.append(
            "Charge accounts do not all have the same number of segments: %s. Every account must follow "
            "the same chart of accounts structure, otherwise segments map to the wrong Oracle segment. "
            "Set 'charge_account_segments' in the mapping file to the correct count and fix the odd rows."
            % detail)
    return errors


def check_splits(rows):
    """Schedule quantities must add up to the line, distributions to the schedule."""
    errors = []
    line_total, scheds, dists = {}, {}, {}

    for rec in rows:
        po = txt(rec.get("po_number"))
        line = txt(as_num(rec.get("line_number")))
        sched = txt(as_num(rec.get("schedule_number")) or 1)
        dist = txt(as_num(rec.get("distribution_number")) or 1)
        goods = txt(rec.get("line_type")) == "Goods"

        qty, price = as_num(rec.get("quantity")), as_num(rec.get("unit_price"))
        amt = as_num(rec.get("line_amount"))
        if goods:
            total = qty
        else:
            total = amt if amt not in ("", None) else \
                (qty * price if qty not in ("", None) and price not in ("", None) else None)
        if isinstance(total, (int, float)):
            line_total.setdefault((po, line), total)

        sv = as_num(rec.get("schedule_quantity"))
        scheds.setdefault((po, line), {})[sched] = (rec["_row"], sv)
        dv = as_num(rec.get("distribution_quantity"))
        dists.setdefault((po, line, sched), {})[dist] = (rec["_row"], dv)

    def total_of(bucket, fallback, what, where):
        vals, blanks = [], []
        for k, (row, v) in sorted(bucket.items()):
            if isinstance(v, (int, float)):
                vals.append(v)
            else:
                blanks.append((k, row))
        if len(bucket) == 1:
            return None                     # single split - nothing to reconcile
        if blanks:
            errors.append("%s: %s %s has no quantity/amount. When a %s is split you must "
                          "enter the quantity or amount for every part."
                          % (", ".join("row %d" % r for _k, r in blanks), what,
                             ", ".join(str(k) for k, _r in blanks), where))
            return None
        return sum(vals)

    for (po, line), bucket in scheds.items():
        got = total_of(bucket, None, "Schedule No.", "line")
        want = line_total.get((po, line))
        if got is not None and isinstance(want, (int, float)) and abs(got - want) > 0.005:
            errors.append("PO %s line %s: the schedule quantities/amounts add up to %s but the "
                          "line total is %s. They must match." % (po, line, got, want))

    for (po, line, sched), bucket in dists.items():
        got = total_of(bucket, None, "Distribution No.", "schedule")
        want = scheds.get((po, line), {}).get(sched, (None, None))[1]
        if want in ("", None):
            want = line_total.get((po, line)) if len(scheds.get((po, line), {})) == 1 else None
        if got is not None and isinstance(want, (int, float)) and abs(got - want) > 0.005:
            errors.append("PO %s line %s schedule %s: the distribution quantities/amounts add up "
                          "to %s but the schedule total is %s. They must match."
                          % (po, line, sched, got, want))
    return errors


# --------------------------------------------------------------- build FBDI
def build_fbdi(rows, batch_id, approval_action, import_source):
    contexts = getattr(P, "DFF_CONTEXTS", {}) or {}
    dff_keys = {}
    for c in P.CATALOGUE:
        if c.get("grp") == "DFF" and c.get("tgt"):
            dff_keys.setdefault(c["tgt"][0], []).append(c["key"])
    """Return {sheet name: [ {fbdi column caption: value}, ... ]} in load order."""
    out = {s: [] for s in SHEET_ORDER}
    h_keys, l_keys, s_keys, d_keys = {}, {}, {}, {}

    def field(rec, sheet):
        # Catalogue, not profile: the row may carry advanced fields, either from
        # the wide sheet or from a client mapping, and they must still be loaded.
        d = {}
        for col in P.CATALOGUE:
            if not col["tgt"] or col["tgt"][0] != sheet:
                continue
            cap = col["tgt"][1]
            if cap.endswith("*"):
                continue                        # charge account, handled separately
            v = rec.get(col["key"])
            if col["dv"] == "date":
                v = as_date(v)
            elif col["dv"] == "num":
                v = as_num(v)
                v = "" if v is None else v
            else:
                v = txt(v)
            if v != "":
                d[cap] = v
        # A DFF context travels with its segments - but ONLY when the row
        # actually carries one. Stamping a context on a row with no attribute
        # values makes Oracle validate that context's required segments against
        # nothing, and the row rejects.
        if sheet in contexts and any(txt(rec.get(k)) for k in dff_keys.get(sheet, ())):
            d["ATTRIBUTE_CATEGORY"] = contexts[sheet]
        return d

    for rec in rows:
        po = txt(rec.get("po_number"))
        line = txt(as_num(rec.get("line_number")))
        sched = txt(as_num(rec.get("schedule_number")) or 1)
        dist = txt(as_num(rec.get("distribution_number")) or 1)

        # ---------------- header
        if po not in h_keys:
            hk = "H%d" % (len(h_keys) + 1)
            h_keys[po] = hk
            d = field(rec, "PO_HEADERS_INTERFACE")
            if rec.get("_auto_po"):
                # The grouping identity is ours, not the client's. Oracle numbers
                # the document itself when Order is blank.
                d.pop("Order", None)
            d["Interface Header Key"] = hk
            d["Action"] = "ORIGINAL"
            d["Batch ID"] = as_num(batch_id) if as_num(batch_id) is not None else batch_id
            d["Ship-to Location"] = txt(rec.get("ship_to_location"))
            if import_source:
                d["Import Source"] = import_source
            if approval_action:
                d["Approval Action"] = approval_action
            out["PO_HEADERS_INTERFACE"].append(d)
        hk = h_keys[po]

        # ---------------- line
        lt = txt(rec.get("line_type"))
        qty = as_num(rec.get("quantity"))
        price = as_num(rec.get("unit_price"))
        amount = as_num(rec.get("line_amount"))
        if amount in ("", None) and qty not in ("", None) and price not in ("", None):
            amount = round(qty * price, 6)

        if (po, line) not in l_keys:
            lk = "%sL%s" % (hk, line)
            l_keys[(po, line)] = lk
            d = field(rec, "PO_LINES_INTERFACE")
            d["Interface Line Key"] = lk
            d["Interface Header Key"] = hk
            d["Action"] = "ADD"
            d["Line"] = as_num(rec.get("line_number"))
            # When a line is split over several rows the line total is the sum
            # of its schedules, worked out in derive_structure.
            line_qty = rec.get("_line_total_qty")
            if line_qty not in (None, ""):
                d["Quantity"] = line_qty
                # Amount has to follow the rolled-up quantity, not the quantity
                # on the first row. A goods line split 250 + 150 at 18.50 is
                # worth 7400, and leaving Amount at the first row's 4625 puts a
                # quantity and an amount on the same line that contradict.
                if price not in ("", None):
                    amount = round(line_qty * price, 6)
            if rec.get("_line_total_amount") not in (None, ""):
                d["Amount"] = rec["_line_total_amount"]
            elif amount not in ("", None):
                d["Amount"] = amount
            if lt != "Goods":
                d.pop("Quantity", None)
                d.pop("Price", None)
                d.pop("UOM", None)
            out["PO_LINES_INTERFACE"].append(d)
        lk = l_keys[(po, line)]

        # ---------------- schedule
        if (po, line, sched) not in s_keys:
            sk = "%sS%s" % (lk, sched)
            s_keys[(po, line, sched)] = sk
            d = field(rec, "PO_LINE_LOCATIONS_INTERFACE")
            d["Interface Line Location Key"] = sk
            d["Interface Line Key"] = lk
            d["Schedule"] = as_num(rec.get("schedule_number")) or 1
            if lt == "Goods":
                d["Quantity"] = as_num(rec.get("schedule_quantity")) or qty
                d.pop("Amount", None)
            else:
                d.pop("Quantity", None)
                d["Amount"] = as_num(rec.get("schedule_quantity")) or amount
            out["PO_LINE_LOCATIONS_INTERFACE"].append(d)
        sk = s_keys[(po, line, sched)]

        # ---------------- distribution
        if (po, line, sched, dist) not in d_keys:
            dk = "%sD%s" % (sk, dist)
            d_keys[(po, line, sched, dist)] = dk
            d = field(rec, "PO_DISTRIBUTIONS_INTERFACE")
            d["Interface Distribution Key"] = dk
            d["Interface Line Location Key"] = sk
            d["Distribution"] = as_num(rec.get("distribution_number")) or 1
            # fall back to the schedule value first, then to the line value, so a
            # split schedule is not silently re-inflated to the full line quantity
            sched_val = as_num(rec.get("schedule_quantity"))
            dist_val = as_num(rec.get("distribution_quantity"))
            if lt == "Goods":
                d["Quantity"] = dist_val or sched_val or qty
                d.pop("Amount", None)
            else:
                d.pop("Quantity", None)
                d["Amount"] = dist_val or sched_val or amount
            acct = txt(rec.get("charge_account"))
            if acct:
                for i, seg in enumerate(re.split(r"[-.|]", acct), start=1):
                    if i <= 30:
                        d["CHARGE_ACCOUNT_SEGMENT%d" % i] = seg.strip()
            out["PO_DISTRIBUTIONS_INTERFACE"].append(d)

    return out


# --------------------------------------------------------------- write outputs
def sheet_layout():
    """{sheet: [captions in Oracle column order]}.

    Read from the baked asset so a conversion never depends on the .xlsm being
    present; falls back to the Oracle template itself if the asset is missing.
    """
    if os.path.exists(LAYOUT_JSON):
        with open(LAYOUT_JSON, encoding="utf-8") as fh:
            return json.load(fh)["sheets"]
    if not os.path.exists(ORACLE_TEMPLATE):
        raise SystemExit("Neither %s nor %s is available - cannot determine the "
                         "Oracle column order." % (LAYOUT_JSON, ORACLE_TEMPLATE))
    wb = openpyxl.load_workbook(ORACLE_TEMPLATE, read_only=True, data_only=True)
    layout = {}
    for s in SHEET_ORDER:
        ws = wb[s]
        for row in ws.iter_rows(min_row=ORA_HEADER_ROW, max_row=ORA_HEADER_ROW):
            layout[s] = [c.value for c in row]
            break
    wb.close()
    return layout


def write_csvs(data, layout, outdir):
    paths = []
    for s in SHEET_ORDER:
        caps = layout[s]
        index = {}
        for i, cap in enumerate(caps):
            index.setdefault(norm(cap), i)
        path = os.path.join(outdir, s + ".csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
            for rec in data[s]:
                line = [""] * len(caps)
                for cap, val in rec.items():
                    i = index.get(norm(cap))
                    if i is None:
                        raise SystemExit("Column '%s' not found on Oracle sheet %s" % (cap, s))
                    line[i] = val
                w.writerow(line)
        paths.append(path)
    zpath = os.path.join(outdir, "PoImportOrders.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
    return paths, zpath


def write_xlsm(data, outdir):
    """Fill a copy of the Oracle template so the mapping can be reviewed / re-generated.

    Delegates to fill_oracle_template, which patches the worksheet XML inside a
    copy of the original .xlsm and copies every other zip part verbatim. Saving
    the workbook through openpyxl instead rewrites xl/drawings/drawing1.xml and
    silently drops the macro-attached "Generate CSV File" button.
    """
    import fill_oracle_template
    return fill_oracle_template.fill(outdir)


# --------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Convert the standard PO template into Oracle FBDI files.")
    ap.add_argument("infile", help="filled client workbook")
    ap.add_argument("--map", dest="mapfile", default=None,
                    help="mapping JSON, for a client file that is not the standard template")
    ap.add_argument("--outdir", default=None, help="output folder (default: output/<file name>)")
    ap.add_argument("--batch", default="1001", help="Batch ID written on every header row")
    ap.add_argument("--approval", default="BYPASS", choices=["BYPASS", "SUBMIT", "NONE", ""],
                    help="PO_HEADERS Approval Action. BYPASS (default) imports the POs already "
                         "approved without running approval rules, which is what an open-PO "
                         "migration wants. SUBMIT sends every PO into the approval workflow. "
                         "NONE / blank leaves them Incomplete.")
    ap.add_argument("--dff", metavar="PATH",
                    help="client descriptive-flexfield spec used to build their "
                         "template. Must be the SAME file, or the extra columns "
                         "are read into the wrong attributes.")
    ap.add_argument("--source", default="", help="Import Source value (optional)")
    ap.add_argument("--xlsm", action="store_true", help="also fill a copy of the Oracle .xlsm template")
    ap.add_argument("--auto-number", dest="auto_number", action="store_true",
                    help="the client leaves PO Number blank and Oracle assigns it from the "
                         "instance's document numbering. Rows are grouped into POs by their "
                         "PO-level fields and the Order column is written blank.")
    ap.add_argument("--derive-ship-to-org", dest="derive_ship_to_org", action="store_true",
                    help="leave Ship-to Organization blank on INVENTORY lines and let Oracle "
                         "resolve it from the Ship-to Location's inventory-organization "
                         "assignment. Downgrades the blank-org error to a warning.")
    ap.add_argument("--force", action="store_true", help="write the CSVs even when validation errors were found")
    a = ap.parse_args()

    outdir = a.outdir or os.path.join(ROOT, "output",
                                      os.path.splitext(os.path.basename(a.infile))[0])
    os.makedirs(outdir, exist_ok=True)

    applied = []
    mapping = None
    if a.mapfile:
        with open(a.mapfile, encoding="utf-8") as fh:
            mapping = json.load(fh)

    # The client's own fields have to be known before the sheet is read, so the
    # extra captions are recognised, and before the FBDI rows are built, so the
    # values land in the right ATTRIBUTE columns.
    dff_notes = []
    dff_path = a.dff or (mapping or {}).get("dff")
    if dff_path:
        import po_dff
        if not os.path.isabs(dff_path) and not os.path.exists(dff_path):
            # A path inside a mapping file is written relative to that mapping
            # file; one typed on the command line is relative to where the
            # command was run. Try both, then the project root.
            bases = []
            if a.mapfile:
                bases.append(os.path.dirname(os.path.abspath(a.mapfile)))
            bases += [os.getcwd(), ROOT]
            tried = []
            for base in bases:
                cand = os.path.normpath(os.path.join(base, dff_path))
                tried.append(cand)
                if os.path.exists(cand):
                    dff_path = cand
                    break
            else:
                raise SystemExit("DFF spec '%s' not found. Looked in: %s"
                                 % (dff_path, ", ".join(tried)))
        with open(LAYOUT_JSON, encoding="utf-8") as fh:
            layout = json.load(fh)
        dff_cols, dff_contexts, dff_notes = po_dff.load(dff_path, layout, P.CATALOGUE)
        P.attach_dff(dff_cols, dff_contexts)

    if a.mapfile:
        rows, read_errors, applied, present = read_with_mapping(a.infile, mapping, a.auto_number)
    else:
        rows, read_errors, present = read_client_sheet(a.infile, a.auto_number)

    expected_segments = (mapping or {}).get("charge_account_segments")
    derived_notes = derive_structure(rows, present)
    errors, warnings = validate(rows, expected_segments, present, a.derive_ship_to_org)
    errors = read_errors + errors

    report = [
        "PO FBDI conversion report",
        "source file : %s" % os.path.abspath(a.infile),
        "layout      : %s" % ("mapping " + os.path.abspath(a.mapfile) if a.mapfile
                              else "standard template"),
        "generated   : %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "data rows   : %d" % len(rows),
        "batch id    : %s" % a.batch,
        "approval    : %s" % (a.approval.strip() or "(blank - POs land Incomplete)"),
        "derived here: %s" % (", ".join(sorted(
            k for k in P.DERIVED_WHEN_ABSENT if k not in present)) or "nothing"),
        "ship-to org : %s" % ("derived by Oracle from Ship-to Location (--derive-ship-to-org)"
                              if a.derive_ship_to_org else "taken from the file"),
        "po numbers  : %s" % ("assigned by Oracle (--auto-number)" if a.auto_number
                              else "taken from the file"),
        "client dff  : %s" % ("; ".join(dff_notes) if dff_notes else "none"),
        "coa segments: %s" % (expected_segments if expected_segments else "not pinned - counts only checked for consistency"),
        "",
    ]
    if AUTO_PO_NOTES:
        report.append("PO NUMBERING - Oracle assigns the document number:")
        report += ["  - " + x for x in AUTO_PO_NOTES]
        report.append("")
    if applied:
        report.append("OVERRIDES APPLIED (%d) - corrections made outside the client file:" % len(applied))
        report += ["  - " + x for x in applied]
        report.append("")
    if derived_notes:
        report.append("SPLIT LINES (%d) - check these were meant to be split:" % len(derived_notes))
        report += ["  - " + n for n in derived_notes]
        report.append("")

    if errors:
        report.append("ERRORS (%d) - these rows will fail in Oracle:" % len(errors))
        report += ["  - " + e for e in errors]
        report.append("")
    if warnings:
        report.append("WARNINGS (%d):" % len(warnings))
        report += ["  - " + w for w in warnings]
        report.append("")

    if errors and not a.force:
        report.append("No CSV files were generated. Fix the errors above, or re-run with --force.")
        path = os.path.join(outdir, "validation_report.txt")
        open(path, "w", encoding="utf-8").write("\n".join(report))
        print("\n".join(report[:40]))
        print("\nfull report: %s" % path)
        sys.exit(1)

    data = build_fbdi(rows, a.batch, a.approval.strip(), a.source.strip())
    layout = sheet_layout()
    paths, zpath = write_csvs(data, layout, outdir)

    report.append("Generated:")
    for s in SHEET_ORDER:
        report.append("  %-32s %d row(s)" % (s + ".csv", len(data[s])))
    report.append("  %-32s upload this file" % "PoImportOrders.zip")
    if a.xlsm:
        report.append("  %-32s for review" % os.path.basename(write_xlsm(data, outdir)))

    path = os.path.join(outdir, "validation_report.txt")
    open(path, "w", encoding="utf-8").write("\n".join(report))
    print("\n".join(report))
    print("\noutput folder: %s" % outdir)


if __name__ == "__main__":
    main()
